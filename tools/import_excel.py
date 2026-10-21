"""Первичная загрузка базы из учётных Excel-файлов филиала.

Запуск:  python -m tools.import_excel "<папка Исходные данные>" [--replace]

Список ОС берётся из ведомости амортизации с отбором по МОЛ, кабинеты — из раскладки
по кабинетам и акта передачи, статус «Предлагается к списанию» — из реестров
списания, ТМЗ — из блоков «ТМЗ» на листах кабинетов. Всё, что не удалось
сопоставить, попадает в отчёт import_report.xlsx.
"""
import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

from . import db

INV_RE = re.compile(r"^2[47]\d{10}$")

ROOM_NAMES = {
    "1": "Служебное помещение", "2": "Учебный кабинет", "3": "Отдел административной работы",
    "5": "Кафедра естественно-математического направления", "6": "Учебный кабинет",
    "7": "Учебный кабинет", "8": "Учебный кабинет", "9": "Библиотека",
    "10": "Отдел организации учебной работы", "11": "Учебный кабинет", "20": "Учебный кабинет",
    "21": "Кабинет", "22": "Учебный кабинет", "23": "Учебный кабинет", "24": "Приёмная",
    "25": "Кафедра общественно-гуманитарного направления", "26": "Өрлеу Hall",
    "27": "Учебный кабинет", "28": "Учебный кабинет",
    "29": "Служба информационных технологий и медиа",
    "Склад 1": "Склад", "Склад 2": "Склад", "Склад 3": "Склад", "Серверная": "Серверная",
    "Под лестницей": "Помещение под лестницей", "Архив": "Архив", "Охрана": "Охрана",
    "Электрощитовая": "Электрощитовая",
}

# Справочник сотрудников (ФИО, подразделение, должность) хранится вне репозитория:
# <папка Исходные данные>/сотрудники.json  — {"ФИО": {"department": "...", "position": "..."}}
STAFF_FILE = "сотрудники.json"


def read_staff(src: Path) -> dict:
    p = src / STAFF_FILE
    if not p.exists():
        return {}
    data = json.loads(p.read_text(encoding="utf-8"))
    return {name: (d.get("department", ""), d.get("position", "")) for name, d in data.items()}

DEPARTMENT_ROOMS = {
    "Руководство": "24", "Кафедра ЕМН": "5", "Кафедра ОГН": "25",
    "Служба информационных технологий и медиа": "29",
    "Отдел организации учебной работы": "10", "Отдел АР": "3",
}

CATEGORY_RULES = [
    ("Программное обеспечение", r"winrar|abbyy|adobe|corel|программное обеспечение|1с|microsoft|windows|антивирус"),
    ("Моноблоки", r"моноблок"),
    ("Ноутбуки", r"ноутбук|notebook|laptop"),
    ("Интерактивные панели", r"интерактивн.*(панел|доск)|панель интерактив|интерактивный кабинет"),
    ("Принтеры и МФУ", r"принтер|мфу|многофункц|сканер|ксерокс|копир|плоттер"),
    ("Мониторы", r"монитор"),
    ("Проекторы", r"проектор|экран"),
    ("Сетевое оборудование", r"коммутатор|маршрутиз|роутер|modem|модем|сервер|wi-?fi|точка доступа|свитч|шкаф|бесперебойн|накопитель на жестк"),
    ("Робототехника", r"робот"),
    ("Аудио и видео", r"микрофон|видео|камер|колонк|акуст|звук|телевизор|фото|регистратор|наушник|пульт"),
    ("Компьютеры", r"компьютер|системный блок|пк\b|рабочая станция"),
]


def categorize(name: str) -> str:
    low = name.lower()
    for category, pattern in CATEGORY_RULES:
        if re.search(pattern, low):
            return category
    return "Прочая техника"


def normalize_room(raw) -> str:
    """'27 кабинет (комплект)' -> '27', 'склад№1' -> 'Склад 1', 'под лестницей' -> 'Под лестницей'."""
    s = str(raw or "").strip()
    if not s:
        return ""
    low = s.lower()
    m = re.search(r"склад\s*№?\s*(\d+)", low)
    if m:
        return f"Склад {m.group(1)}"
    if "лестниц" in low:
        return "Под лестницей"
    if "сервер" in low:
        return "Серверная"
    # Только однозначные записи: «Монитор в 27, процессор в 29» кабинет не определяет.
    m = re.match(r"(?:каб(?:инет)?\.?\s*№?\s*)?(\d{1,3})\b(?!.*\d)", low)
    if m:
        return m.group(1)
    return ""


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def rows(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            yield ws.title, [cell_str(c) for c in row]


def find_file(src: Path, name: str) -> Path | None:
    hits = [p for p in src.rglob(name) if "_Дубликаты" not in p.parts]
    return hits[0] if hits else None


def find_mol_file(src: Path) -> Path | None:
    """Ведомость амортизации с отбором «МОЛ Равно …» и наибольшим числом строк."""
    best, best_n = None, 0
    for p in sorted(src.rglob("*.xlsx")):
        if "_Дубликаты" in p.parts:
            continue
        try:
            items, mol = read_mol_list(p)
        except Exception:
            continue
        if mol and len(items) > best_n:
            best, best_n = p, len(items)
    return best


def to_float(s: str):
    try:
        return float(str(s).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def read_mol_list(path: Path):
    """Ведомость амортизации по МОЛ: № | Инв. номер | Актив | Дата | ... | Первонач. | ... | Остаточная."""
    items, mol, col = {}, "", {}
    for _, r in rows(path):
        m = re.search(r'МОЛ Равно\s*"([^"]+)"', " ".join(r))
        if m:
            mol = m.group(1)
        if "Инв. номер" in r:
            col = {h: i for i, h in enumerate(r) if h}
            continue
        if not col:
            continue
        inv = r[col["Инв. номер"]]
        if not INV_RE.match(inv):
            continue
        get = lambda h: r[col[h]] if h in col and col[h] < len(r) else ""
        items[inv] = {
            "name": get("Долгосрочный актив"),
            "acquired": get("Дата поступления"),
            "cost": to_float(get("Первоначальная стоимость")),
            "residual": to_float(get("Остаточная стоимость")),
        }
    return items, mol


def read_locations(path: Path):
    """Любой лист, где в строке есть инв. номер и ячейка-местонахождение (последняя текстовая)."""
    locs = {}
    for sheet, r in rows(path):
        invs = [c for c in r if INV_RE.match(c)]
        if not invs:
            continue
        place = ""
        for c in reversed(r):
            if c and not INV_RE.match(c) and re.search(r"кабинет|склад|лестниц|сервер|каб\.", c.lower()):
                place = c
                break
        room = normalize_room(place) or (normalize_room(sheet) if re.match(r"^\d+$|склад|лестниц", sheet.lower()) else "")
        locs[invs[0]] = (room, place)
    return locs


def read_room_stock(path: Path):
    """Блоки «ТМЗ» внизу листов кабинетов: наименование | количество | примечание."""
    out = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        in_block = False
        for row in ws.iter_rows(values_only=True):
            cells = [cell_str(c) for c in row if cell_str(c)]
            if not cells:
                continue
            if cells[0].upper() == "ТМЗ":
                in_block = True
                continue
            if in_block and len(cells) >= 2:
                m = re.match(r"^\s*([\d.,]+)\s*(.*)$", cells[1])
                qty = to_float(m.group(1)) if m else None
                if qty is None:
                    continue
                unit = (m.group(2) or "шт.").strip() or "шт."
                unit = {"шт": "шт.", "м": "м"}.get(unit, unit)
                out.append({
                    "name": re.sub(r"\s+", " ", cells[0]).strip(),
                    "qty": qty, "unit": unit,
                    "room": normalize_room(ws.title),
                    "notes": cells[2] if len(cells) > 2 else "",
                })
    return out


def read_inv_numbers(path: Path) -> set[str]:
    return {c for _, r in rows(path) for c in r if INV_RE.match(c)}


def run(src: Path, conn, replace: bool = False, report_path: Path | None = None) -> dict:
    db.init(conn)
    if conn.execute("SELECT COUNT(*) FROM assets").fetchone()[0] and not replace:
        raise SystemExit("В базе уже есть данные. Для полной перезагрузки добавьте --replace (текущая база будет стерта).")

    mol_file = find_mol_file(src)
    if not mol_file:
        raise SystemExit(f"Не найдена ведомость амортизации с отбором по МОЛ в {src}")
    assets, mol = read_mol_list(mol_file)

    loc_sources = [
        ("Оргтехника и комплектующие по кабинетам.xlsx", "по кабинетам"),
        ("Нуржан (оргтехника по кабинетам).xlsx", "акт передачи 23.01.2026"),
    ]
    room_of, conflicts, place_notes = {}, [], defaultdict(list)
    for fname, label in loc_sources:
        p = find_file(src, fname)
        if not p:
            continue
        for inv, (room, place) in read_locations(p).items():
            if place and normalize_room(place) != room or (place and not room):
                place_notes[inv].append(f"{label}: {place}")
            if not room:
                continue
            if inv in room_of and room_of[inv][0] != room:
                conflicts.append((inv, room_of[inv][0], room_of[inv][1], room, label))
                continue
            room_of.setdefault(inv, (room, label))

    writeoff = set()
    for fname in ["Оргтехника на списание 2025.xlsx", "Реестр машины и оборудования.xlsx", "Списание 2026.xlsx"]:
        p = find_file(src, fname)
        if p:
            writeoff |= read_inv_numbers(p)

    stock = []
    p = find_file(src, "Оргтехника и комплектующие по кабинетам.xlsx")
    if p:
        stock = read_room_stock(p)

    ts = db.now()
    with db.tx(conn):
        if replace:
            for t in ["audit_scans", "audits", "movements", "stock_moves", "stock_items", "assets", "employees", "rooms", "events"]:
                conn.execute(f"DELETE FROM {t}")
        codes = set(ROOM_NAMES) | {r for r, _ in room_of.values()} | {s["room"] for s in stock if s["room"]}
        for i, code in enumerate(sorted(codes, key=db.room_code_key)):
            conn.execute("INSERT INTO rooms (code, name, sort) VALUES (?, ?, ?)", (code, ROOM_NAMES.get(code, ""), i))
        room_id = {r["code"]: r["id"] for r in conn.execute("SELECT id, code FROM rooms")}

        staff = read_staff(src)
        for name, (dept, pos) in staff.items():
            conn.execute(
                "INSERT INTO employees (full_name, department, position, room_id) VALUES (?, ?, ?, ?)",
                (name, dept, pos, room_id.get(DEPARTMENT_ROOMS.get(dept, ""))),
            )

        for inv, a in sorted(assets.items()):
            room = room_of.get(inv, ("", ""))[0]
            status = "Предлагается к списанию" if inv in writeoff else "В эксплуатации"
            if room.startswith("Склад") and status == "В эксплуатации":
                status = "На складе"
            conn.execute(
                """INSERT INTO assets (inv, name, category, room_id, mol, condition, status, acquired,
                                       cost, residual, notes, source, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, 'Не проверено', ?, ?, ?, ?, ?, ?, ?, ?)""",
                (inv, a["name"], categorize(a["name"]), room_id.get(room), mol, status, a["acquired"],
                 a["cost"], a["residual"], "; ".join(place_notes.get(inv, [])), mol_file.name, ts, ts),
            )

        for s in stock:
            conn.execute(
                "INSERT INTO stock_items (name, unit, qty, min_qty, room_id, notes) VALUES (?, ?, ?, 0, ?, ?)",
                (s["name"], s["unit"], s["qty"], room_id.get(s["room"]), s["notes"]),
            )
        db.log_event(conn, "импорт", "Первичная загрузка", mol_file.name,
                     f"ОС: {len(assets)}, с кабинетом: {sum(1 for i in assets if i in room_of)}, ТМЗ: {len(stock)}")

    no_room = [inv for inv in assets if inv not in room_of]
    not_in_mol = sorted(set(room_of) - set(assets))
    writeoff_missing = sorted(writeoff - set(assets))
    summary = {
        "assets": len(assets), "with_room": len(assets) - len(no_room), "no_room": len(no_room),
        "writeoff": len(writeoff & set(assets)), "stock": len(stock), "conflicts": len(conflicts),
        "rooms": len(room_id), "employees": len(staff),
    }
    if report_path:
        write_report(report_path, assets, no_room, conflicts, not_in_mol, room_of, writeoff_missing)
    return summary


def write_report(path, assets, no_room, conflicts, not_in_mol, room_of, writeoff_missing):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Без кабинета"
    ws.append(["Инв. номер", "Наименование", "Что сделать"])
    for inv in sorted(no_room):
        ws.append([inv, assets[inv]["name"], "Найти при инвентаризации и указать кабинет"])
    ws = wb.create_sheet("Разные кабинеты")
    ws.append(["Инв. номер", "Наименование", "Кабинет (принят)", "Источник", "Кабинет (другой файл)", "Источник"])
    for inv, r1, s1, r2, s2 in conflicts:
        ws.append([inv, assets.get(inv, {}).get("name", ""), r1, s1, r2, s2])
    ws = wb.create_sheet("Нет в списке МОЛ")
    ws.append(["Инв. номер", "Кабинет по раскладке", "Комментарий"])
    for inv in not_in_mol:
        ws.append([inv, room_of[inv][0], "Есть в раскладке по кабинетам, но нет в ведомости МОЛ — уточнить в бухгалтерии"])
    ws = wb.create_sheet("Списание вне списка МОЛ")
    ws.append(["Инв. номер", "Комментарий"])
    for inv in writeoff_missing:
        ws.append([inv, "Есть в реестрах списания, но не числится за МОЛ"])
    for sheet in wb.worksheets:
        for col in "ABCDEF":
            sheet.column_dimensions[col].width = 28
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Загрузка базы ӨRLEU BASE из Excel")
    ap.add_argument("source", type=Path, help="Папка с исходными Excel-файлами")
    ap.add_argument("--replace", action="store_true", help="Стереть текущую базу и загрузить заново")
    ap.add_argument("--db", type=Path, default=db.DB_PATH)
    args = ap.parse_args(argv)
    conn = db.connect(args.db)
    report = args.db.parent / "import_report.xlsx"
    summary = run(args.source, conn, args.replace, report)
    print("Загружено:")
    for k, label in [("assets", "ОС"), ("with_room", "с кабинетом"), ("no_room", "без кабинета"),
                     ("writeoff", "предлагается к списанию"), ("stock", "позиций ТМЗ"),
                     ("rooms", "кабинетов"), ("employees", "сотрудников"), ("conflicts", "расхождений кабинетов")]:
        print(f"  {label}: {summary[k]}")
    print(f"Отчёт расхождений: {report}")


if __name__ == "__main__":
    sys.exit(main())
